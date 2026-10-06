"""Tests de picos, colores, base de datos, exportación a Origin e idiomas."""
import io
import re
import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from cvcap import colors, peaks as pk, report
from cvcap.analysis import ANODIC, CATHODIC, analyze
from cvcap.db import Database, record_from_result
from cvcap.i18n import TEXTS, CvcapError, t
from cvcap.io import CVData

ROOT = Path(__file__).resolve().parents[1]


def graphite_like_cv(cycles=3, rate=1e-3, step=1e-3):
    """CV con un pico SEI (0,65 V, solo ciclo 1) y un pico anódico cuyo potencial crece con el ciclo."""
    legs = [np.arange(2.0, 0.01, -step)]
    for c in range(cycles):
        legs.append(np.arange(0.01, 2.0, step))
        if c < cycles - 1:
            legs.append(np.arange(2.0, 0.01, -step))
    e = np.concatenate(legs)
    tt = np.arange(len(e)) * step / rate
    direction = np.sign(np.gradient(e))
    cycle = np.cumsum(np.r_[0, (np.diff(direction) < 0).astype(int)]) + 1
    i = np.zeros_like(e)
    cat = direction < 0
    i[cat] -= 1e-3 * np.exp(-(e[cat] - 0.01) / 0.07)
    first = cat & (cycle == 1)
    i[first] -= 2e-4 * np.exp(-((e[first] - 0.65) / 0.05) ** 2)
    an_mask = ~cat
    centre = 0.28 + 0.01 * (cycle[an_mask] - 1)
    i[an_mask] += 8e-4 * np.exp(-((e[an_mask] - centre) / 0.05) ** 2)
    return CVData(potential_V=e, current_A=i, time_s=tt)


class TestPeaks(unittest.TestCase):
    def setUp(self):
        self.res = analyze(graphite_like_cv(), mass_mg=2.0)
        self.windows = [pk.PeakWindow("SEI", CATHODIC, 0.4, 1.0), pk.PeakWindow("Delit", ANODIC, 0.1, 0.5)]

    def test_window_peaks_positions(self):
        df = pk.window_peaks(self.res, self.windows)
        sei1 = df[(df["name"] == "SEI") & (df["cycle"] == 1)].iloc[0]
        self.assertAlmostEqual(sei1["e_peak"], 0.65, delta=0.01)
        self.assertFalse(sei1["edge"])
        sei2 = df[(df["name"] == "SEI") & (df["cycle"] == 2)].iloc[0]
        self.assertTrue(sei2["edge"])  # sin pico SEI tras el ciclo 1: el extremo cae en el borde
        delit = df[df["name"] == "Delit"].sort_values("cycle")["e_peak"].to_numpy()
        np.testing.assert_allclose(delit, [0.28, 0.29, 0.30], atol=0.006)

    def test_reference_shift_and_separation(self):
        df = pk.add_reference_shifts(pk.window_peaks(self.res, self.windows), 1)
        shift3 = df[(df["name"] == "Delit") & (df["cycle"] == 3)]["de_ref"].iloc[0]
        self.assertAlmostEqual(shift3, 20.0, delta=6.0)
        sep = pk.peak_separation(df, "Delit", "SEI")
        self.assertEqual(len(sep), 3)
        self.assertAlmostEqual(sep["dep"].iloc[0], (0.28 - 0.65) * 1000, delta=15)

    def test_auto_peaks(self):
        auto = pk.auto_peaks(self.res, prominence=0.05)
        anodic = auto[auto["kind"] == ANODIC]
        self.assertEqual(sorted(anodic["cycle"].unique().tolist()), [1, 2, 3])
        cath1 = auto[(auto["kind"] == CATHODIC) & (auto["cycle"] == 1)]
        self.assertTrue(((cath1["e_peak"] - 0.65).abs() < 0.02).any())

    def test_windows_from_records_validation(self):
        recs = [{"name": "ok", "kind": "anodic", "e_min": 0.5, "e_max": 0.1},
                {"name": "", "kind": "anodic", "e_min": 0, "e_max": 1},
                {"name": "bad", "kind": None, "e_min": 0, "e_max": 1},
                {"name": "nan", "kind": "cathodic", "e_min": None, "e_max": 1}]
        out = pk.windows_from_records(recs)
        self.assertEqual(len(out), 1)
        self.assertEqual((out[0].e_min, out[0].e_max), (0.1, 0.5))


class TestColors(unittest.TestCase):
    def test_shades(self):
        for ramp in colors.RAMPS:
            shades = colors.cycle_shades(ramp, 5)
            self.assertEqual(len(shades), 5)
            self.assertEqual(shades[0], colors.RAMPS[ramp][0])
            self.assertEqual(shades[-1], colors.RAMPS[ramp][-1])
            self.assertTrue(all(re.fullmatch(r"#[0-9a-f]{6}", c) for c in shades))
        self.assertEqual(colors.cycle_shades("blues", 3, darkest_first=True)[0], colors.RAMPS["blues"][-1])
        self.assertEqual(len(colors.cycle_shades("reds", 1)), 1)
        self.assertEqual(colors.cycle_shades("reds", 0), [])


class TestDatabase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "test.sqlite")
        self.res = analyze(graphite_like_cv(), mass_mg=2.0, name="muestra.txt")
        self.peaks = pk.window_peaks(self.res, pk.default_windows("es"))
        self.meta = {"name": "Grafito A", "material": "commercial", "batch": "L1", "measured_on": "2026-10-05",
                     "prep": {"composition": "80:10:10", "thickness_um": 30.0, "pressing": ""},
                     "notes": "sin prensar"}

    def tearDown(self):
        self.tmp.cleanup()

    def test_roundtrip(self):
        sid = self.db.next_id()
        self.assertEqual(sid, "S-0001")
        rec = record_from_result(self.res, sid, self.meta, peaks=self.peaks)
        self.db.save(rec)
        self.assertEqual(self.db.next_id(), "S-0002")
        got = self.db.get(sid)
        self.assertEqual(got.name, "Grafito A")
        self.assertEqual(got.prep, {"composition": "80:10:10", "thickness_um": 30.0})
        pd.testing.assert_series_equal(got.cycles["cap_delit"], self.res.cycles["cap_delit"], check_names=False)
        np.testing.assert_allclose(got.curve["potential_V"], self.res.cv.potential_V, atol=1e-6)
        np.testing.assert_array_equal(got.curve["point_cycle"], self.res.point_cycle)
        self.assertEqual(len(got.peaks), len(self.peaks))
        self.assertEqual(got.cycle_notes, self.res.cycle_notes)
        catalog = self.db.list_samples()
        self.assertEqual(catalog["id"].tolist(), [sid])
        self.assertAlmostEqual(catalog["ce1"].iloc[0], self.res.summary["ce1"])

    def test_duplicate_and_overwrite(self):
        rec = record_from_result(self.res, "G-01", self.meta)
        self.db.save(rec)
        with self.assertRaises(CvcapError) as ctx:
            self.db.save(record_from_result(self.res, "G-01", self.meta))
        self.assertEqual(ctx.exception.code, "err.db_exists")
        self.db.save(record_from_result(self.res, "G-01", {**self.meta, "name": "B"}), overwrite=True)
        self.assertEqual(self.db.get("G-01").name, "B")

    def test_bad_id(self):
        with self.assertRaises(CvcapError):
            self.db.save(record_from_result(self.res, "   ", self.meta))

    def test_update_delete(self):
        self.db.save(record_from_result(self.res, "S-0001", self.meta))
        self.db.update_metadata("S-0001", notes="prensado 1 t", prep={"pressing": "1 t"})
        got = self.db.get("S-0001")
        self.assertEqual((got.notes, got.prep), ("prensado 1 t", {"pressing": "1 t"}))
        self.db.delete("S-0001")
        self.assertEqual(self.db.count(), 0)
        with self.assertRaises(CvcapError):
            self.db.get("S-0001")

    def test_backup_restore(self):
        self.db.save(record_from_result(self.res, "S-0001", self.meta))
        data = self.db.export_bytes()
        other = Database(Path(self.tmp.name) / "other.sqlite")
        self.assertEqual(other.import_bytes(data), (1, 0))
        self.assertEqual(other.import_bytes(data), (0, 1))
        self.assertEqual(other.import_bytes(data, overwrite=True), (1, 0))
        with self.assertRaises(CvcapError):
            other.import_bytes(b"not a database")


class TestOriginExport(unittest.TestCase):
    def test_workbook_layout(self):
        res = analyze(graphite_like_cv(), mass_mg=2.0, name="A")
        res2 = analyze(graphite_like_cv(cycles=2), mass_mg=2.0, name="B")
        p = pk.window_peaks(res, pk.default_windows("en"))
        items = [report.item_from_result(res, "A", peaks=p), report.item_from_result(res2, "B")]
        wb = load_workbook(io.BytesIO(report.origin_workbook(items, "en", "uA")))
        self.assertEqual(wb.sheetnames, ["A CV", "A Cap", "A Peaks", "B CV", "B Cap", "Comparison", "Info"])
        cv = wb["A CV"]
        self.assertEqual(cv.max_column, 6)  # 3 ciclos x (E, I)
        self.assertEqual([cv.cell(1, j).value for j in (1, 2)], ["Potential", "Current"])
        self.assertEqual([cv.cell(2, j).value for j in (1, 2)], ["V", "µA"])
        self.assertEqual(cv.cell(3, 3).value, "A · Cycle 2")
        n1 = int((res.point_cycle == 1).sum())
        self.assertEqual(cv.cell(4, 1).value, res.cv.potential_V[res.point_cycle == 1][0])
        self.assertAlmostEqual(cv.cell(4, 2).value, res.current_A[res.point_cycle == 1][0] * 1e6)
        self.assertIsNotNone(cv.cell(3 + n1, 1).value)
        cap = wb["A Cap"]
        self.assertEqual(cap.cell(2, 3).value, "mAh/g")
        self.assertAlmostEqual(cap.cell(4, 3).value, res.cycles["cap_delit"].iloc[0])

    def test_specific_current(self):
        res = analyze(graphite_like_cv(), mass_mg=2.0, name="A")
        wb = load_workbook(io.BytesIO(report.origin_workbook([report.item_from_result(res)], "es", "A/g")))
        ws = wb.worksheets[0]
        self.assertEqual(ws.cell(1, 2).value, "Corriente específica")
        self.assertAlmostEqual(ws.cell(4, 2).value, res.current_A[res.point_cycle == 1][0] / 0.002)


class TestI18n(unittest.TestCase):
    def test_same_keys_in_both_languages(self):
        self.assertEqual(set(TEXTS["es"]), set(TEXTS["en"]))

    def test_all_keys_used_in_code_exist(self):
        sources = list((ROOT / "cvcap").glob("*.py")) + [ROOT / "app.py"]
        literal = re.compile(r"""\b(?:T|t|Note|CvcapError|NovaFormatError)\(\s*["']([a-z_]+\.[A-Za-z0-9_.]+|[a-z_]+)["']\s*[,)]""")
        missing = set()
        for path in sources:
            for key in literal.findall(path.read_text(encoding="utf-8")):
                if key not in TEXTS["es"]:
                    missing.add(key)
        self.assertEqual(missing, set())

    def test_dynamic_key_families(self):
        from cvcap.analysis import CYCLE_COLUMNS, SEGMENT_COLUMNS, SUMMARY_KEYS
        from cvcap.db import MATERIALS, PREP_FIELDS
        from cvcap.io import ROLES

        keys = [f"col.{c}" for c in CYCLE_COLUMNS + ["notes", "file"]]
        keys += [f"seg.{c}" for c in SEGMENT_COLUMNS] + [f"sum.{k}" for k in SUMMARY_KEYS]
        keys += [f"material.{m}" for m in MATERIALS] + [f"prep.{p}" for p in PREP_FIELDS]
        keys += [f"ph.{p}" for p in PREP_FIELDS] + [f"role.{r}" for r in ROLES]
        keys += [f"ramp.{r}" for r in colors.RAMPS] + [f"kind.{k}" for k in (CATHODIC, ANODIC)]
        keys += [f"kind_short.{k}" for k in (CATHODIC, ANODIC)] + [f"legend.peaks_{k}" for k in (CATHODIC, ANODIC)]
        keys += [f"method.{m}" for m in ("signo", "direccion")] + [f"method_long.{m}" for m in ("signo", "direccion")]
        keys += [f"nav.{p}" for p in ("page", "analyze", "database")] + ["sb.mass_direct", "sb.mass_disc"]
        keys += ["note.sorted_by_time", "note.sorted_by_index"]
        keys += [f"cat.{c}" for c in ("id", "name", "material", "batch", "measured_on", "source_file", "mass_mg",
                                       "method", "n_cycles", "ce1", "ce_mean", "cap_delit1", "cap_delit_last",
                                       "ret_last", "notes", "updated_at")]
        keys += [f"meta.{k}" for k in ("name", "material", "batch", "measured_on", "source_file")]
        self.assertEqual([k for k in keys if k not in TEXTS["es"]], [])

    def test_placeholders_match(self):
        field = re.compile(r"\{(\w+)")
        for key, es in TEXTS["es"].items():
            self.assertEqual(set(field.findall(es)), set(field.findall(TEXTS["en"][key])), key)

    def test_translation(self):
        self.assertEqual(t("cycle_n", "en", n=3), "Cycle 3")
        self.assertEqual(t("cycle_n", "es", n=3), "Ciclo 3")
        self.assertEqual(t("missing.key", "en"), "missing.key")


if __name__ == "__main__":
    unittest.main()
