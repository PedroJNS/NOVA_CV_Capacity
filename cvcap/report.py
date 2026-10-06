"""Tablas resumen y exportaciones (Excel de resultados, CSV y Excel para OriginLab).

Formato OriginLab
-----------------
Cada hoja tiene tres filas de cabecera antes de los datos:

1. *Long Name* (magnitud)   2. *Units* (unidades)   3. *Comments* (muestra · ciclo)

En Origin: *Data → Import from File → Excel (XLS, XLSX, XLSM)* y en el asistente
indique que la fila 1 es Long Name, la 2 Units y la 3 Comments (o use
*Import Wizard → Header lines: 3*). Las columnas van por pares X/Y, un par por
ciclo, listas para *Plot → Line*.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .analysis import AnalysisResult, cycle_label, localize_cycles, localize_summary
from .i18n import Note, t

CURRENT_UNITS = {"mA": 1e3, "uA": 1e6, "A": 1.0, "A/g": None}


@dataclass
class ExportItem:
    """Datos de una muestra listos para exportar (desde un análisis o desde la base de datos)."""

    label: str
    potential_V: np.ndarray
    current_A: np.ndarray
    point_cycle: np.ndarray
    cycles: pd.DataFrame
    cycle_notes: Dict[int, List[Note]] = field(default_factory=dict)
    summary: Dict[str, object] = field(default_factory=dict)
    mass_mg: Optional[float] = None
    peaks: Optional[pd.DataFrame] = None
    meta: Dict[str, object] = field(default_factory=dict)


def item_from_result(result: AnalysisResult, label: Optional[str] = None, peaks=None, meta=None) -> ExportItem:
    return ExportItem(
        label=label or result.name, potential_V=result.cv.potential_V, current_A=result.current_A,
        point_cycle=result.point_cycle, cycles=result.cycles, cycle_notes=result.cycle_notes,
        summary=result.summary, mass_mg=result.mass_mg, peaks=peaks, meta=dict(meta or {}),
    )


def item_from_record(rec) -> ExportItem:
    curve = rec.curve or {}
    empty = np.array([])
    meta = {"name": rec.name, "material": rec.material, "batch": rec.batch, "measured_on": rec.measured_on,
            "source_file": rec.source_file, "area_cm2": rec.area_cm2, "notes": rec.notes, "prep": rec.prep}
    return ExportItem(
        label=rec.id, potential_V=curve.get("potential_V", empty), current_A=curve.get("current_A", empty),
        point_cycle=curve.get("point_cycle", empty.astype(int)), cycles=rec.cycles, cycle_notes=rec.cycle_notes,
        summary=rec.summary, mass_mg=rec.mass_mg, peaks=rec.peaks, meta=meta,
    )


# --------------------------------------------------------------------------
# Tablas
# --------------------------------------------------------------------------
def summary_table(results: List[AnalysisResult], lang: str = "es") -> pd.DataFrame:
    return pd.DataFrame([r.summary_row(lang) for r in results])


def cycles_long_table(results: List[AnalysisResult], lang: str = "es") -> pd.DataFrame:
    frames = []
    for r in results:
        df = r.cycles_table(lang)
        df.insert(0, t("col.file", lang), r.name)
        frames.append(df)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def _sheet_name(base: str, used: set) -> str:
    name = re.sub(r"[\[\]\:\*\?\/\\]", "_", str(base)).strip()[:31] or "Sheet"
    candidate, k = name, 1
    while candidate in used:
        suffix = f"_{k}"
        candidate = name[: 31 - len(suffix)] + suffix
        k += 1
    used.add(candidate)
    return candidate


def results_to_excel(results: List[AnalysisResult], lang: str = "es") -> bytes:
    """Libro con resumen, todos los ciclos y el detalle de segmentos de cada archivo."""
    buffer = io.BytesIO()
    used: set = set()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary_table(results, lang).to_excel(writer, sheet_name=_sheet_name(t("xl.summary", lang), used), index=False)
        cycles_long_table(results, lang).to_excel(writer, sheet_name=_sheet_name(t("xl.cycles", lang), used),
                                                  index=False)
        for r in results:
            stem = re.sub(r"\.[A-Za-z0-9]+$", "", r.name)[:20]
            r.segments_table(lang).to_excel(writer, sheet_name=_sheet_name(f"{stem}_{t('xl.seg', lang)}", used),
                                            index=False)
            if r.notes:
                pd.DataFrame({t("col.notes", lang): r.notes_text(lang)}).to_excel(
                    writer, sheet_name=_sheet_name(f"{stem}_{t('xl.notes', lang)}", used), index=False)
    return buffer.getvalue()


def results_to_csv(results: List[AnalysisResult], lang: str = "es") -> bytes:
    """CSV de todos los ciclos (en español: separador ';' y decimal ','; en inglés: ',' y '.')."""
    table = cycles_long_table(results, lang)
    if lang == "es":
        return table.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")
    return table.to_csv(index=False).encode("utf-8-sig")


# --------------------------------------------------------------------------
# OriginLab
# --------------------------------------------------------------------------
def _write_columns(ws, columns: List[dict]) -> None:
    """columns: [{'long': str, 'units': str, 'comment': str, 'values': array}]"""
    bold = Font(bold=True)
    for j, col in enumerate(columns, start=1):
        ws.cell(row=1, column=j, value=col["long"]).font = bold
        ws.cell(row=2, column=j, value=col["units"])
        ws.cell(row=3, column=j, value=col["comment"])
        for i, v in enumerate(col["values"], start=4):
            if v is None:
                continue
            if isinstance(v, (float, np.floating)) and not np.isfinite(v):
                continue
            ws.cell(row=i, column=j, value=float(v) if isinstance(v, (np.floating, np.integer, float, int))
                    and not isinstance(v, bool) else v)
        ws.column_dimensions[get_column_letter(j)].width = 16


def _current_factor(unit: str, mass_mg: Optional[float]):
    """Devuelve (factor, unidad mostrada). Para A/g sin masa se recurre a mA."""
    if unit == "A/g":
        if mass_mg:
            return 1.0 / (mass_mg / 1000.0), "A/g"
        return 1e3, "mA"
    factor = CURRENT_UNITS.get(unit, 1e3)
    return factor, ("µA" if unit == "uA" else unit)


def origin_workbook(items: List[ExportItem], lang: str = "es", current_unit: str = "mA") -> bytes:
    """Excel con las CV por ciclo (pares E/I en columnas) y las capacidades, listo para OriginLab."""
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()

    for item in items:
        factor, unit = _current_factor(current_unit, item.mass_mg)
        i_long = t("origin.specific_current", lang) if unit == "A/g" else t("origin.current", lang)

        # --- CV por ciclo
        ws = wb.create_sheet(_sheet_name(f"{item.label[:24]} CV", used))
        cols = []
        cycles = sorted(int(c) for c in np.unique(item.point_cycle) if c > 0)
        for c in cycles:
            mask = item.point_cycle == c
            comment = f"{item.label} · {cycle_label(c, lang)}"
            cols.append({"long": t("origin.potential", lang), "units": "V", "comment": comment,
                         "values": item.potential_V[mask]})
            cols.append({"long": i_long, "units": unit, "comment": comment,
                         "values": item.current_A[mask] * factor})
        _write_columns(ws, cols)

        # --- capacidades por ciclo
        ws = wb.create_sheet(_sheet_name(f"{item.label[:24]} {t('origin.cap_sheet', lang)}", used))
        df = item.cycles
        has_mass = "cap_delit" in df.columns
        cap_unit = "mAh/g" if has_mass else "mAh"
        lit, delit = ("cap_lit", "cap_delit") if has_mass else ("mah_lit", "mah_delit")
        cols = [
            {"long": t("col.cycle", lang), "units": "", "comment": item.label, "values": df["cycle"].to_numpy()},
            {"long": t("origin.cap_lit", lang), "units": cap_unit, "comment": item.label, "values": df[lit].to_numpy()},
            {"long": t("origin.cap_delit", lang), "units": cap_unit, "comment": item.label,
             "values": df[delit].to_numpy()},
            {"long": t("origin.ce", lang), "units": "%", "comment": item.label, "values": df["ce"].to_numpy()},
            {"long": t("origin.ret", lang), "units": "%", "comment": item.label, "values": df["ret"].to_numpy()},
        ]
        if "irr" in df.columns:
            cols.append({"long": t("origin.irr", lang), "units": "mAh/g", "comment": item.label,
                         "values": df["irr"].to_numpy()})
        _write_columns(ws, cols)

        # --- picos
        if item.peaks is not None and not item.peaks.empty:
            ws = wb.create_sheet(_sheet_name(f"{item.label[:24]} {t('origin.peaks_sheet', lang)}", used))
            cols = []
            for name, group in item.peaks.groupby("name", sort=False):
                group = group.sort_values("cycle").copy()
                if "edge" in group:  # un extremo en el borde de la ventana no es un pico: celda vacía
                    edge = group["edge"].astype(bool).to_numpy()
                    group.loc[edge, ["e_peak", "i_peak"]] = np.nan
                comment = f"{item.label} · {name}"
                cols.append({"long": t("col.cycle", lang), "units": "", "comment": comment,
                             "values": group["cycle"].to_numpy()})
                cols.append({"long": t("origin.e_peak", lang), "units": "V", "comment": comment,
                             "values": group["e_peak"].to_numpy()})
                cols.append({"long": t("origin.i_peak", lang), "units": unit, "comment": comment,
                             "values": group["i_peak"].to_numpy() * factor})
            _write_columns(ws, cols)

    # --- comparación de muestras
    if len(items) > 1:
        ws = wb.create_sheet(_sheet_name(t("origin.compare_sheet", lang), used))
        cols = []
        for item in items:
            df = item.cycles
            delit = "cap_delit" if "cap_delit" in df.columns else "mah_delit"
            unit = "mAh/g" if delit == "cap_delit" else "mAh"
            cols.append({"long": t("col.cycle", lang), "units": "", "comment": item.label,
                         "values": df["cycle"].to_numpy()})
            cols.append({"long": t("origin.cap_delit", lang), "units": unit, "comment": item.label,
                         "values": df[delit].to_numpy()})
            cols.append({"long": t("origin.ce", lang), "units": "%", "comment": item.label,
                         "values": df["ce"].to_numpy()})
        _write_columns(ws, cols)

    # --- información de las muestras
    ws = wb.create_sheet(_sheet_name(t("origin.info_sheet", lang), used))
    row = 1
    bold = Font(bold=True)
    for item in items:
        ws.cell(row=row, column=1, value=item.label).font = bold
        row += 1
        meta = dict(item.meta)
        prep = meta.pop("prep", {}) or {}
        pairs = []
        for key in ("name", "material", "batch", "measured_on", "source_file"):
            value = meta.get(key)
            if value:
                if key == "material":
                    value = t(f"material.{value}", lang)
                pairs.append((t(f"meta.{key}", lang), value))
        if meta.get("area_cm2"):
            pairs.append((t("meta.area_cm2", lang), meta["area_cm2"]))
        for key, value in prep.items():
            pairs.append((t(f"prep.{key}", lang), value))
        if meta.get("notes"):
            pairs.append((t("meta.notes", lang), meta["notes"]))
        for key, value in localize_summary(item.summary, lang).items():
            pairs.append((key, value))
        for key, value in pairs:
            ws.cell(row=row, column=1, value=key)
            ws.cell(row=row, column=2, value=value)
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def catalog_to_excel(catalog: pd.DataFrame, lang: str = "es") -> bytes:
    """Catálogo de la base de datos (una fila por muestra, con preparación y notas)."""
    df = localize_catalog(catalog, lang)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=t("xl.catalog", lang)[:31], index=False)
    return buffer.getvalue()


def localize_catalog(catalog: pd.DataFrame, lang: str = "es") -> pd.DataFrame:
    df = catalog.copy()
    if df.empty:
        return df
    if "material" in df:
        df["material"] = df["material"].map(lambda m: t(f"material.{m}", lang) if m else "")
    if "method" in df:
        df["method"] = df["method"].map(lambda m: t(f"method.{m}", lang) if m else "")
    prep = None
    if "prep" in df:
        prep = pd.DataFrame(list(df["prep"].map(lambda p: p or {})), index=df.index)
        df = df.drop(columns=["prep"])
    df = df.rename(columns={c: t(f"cat.{c}", lang) for c in df.columns})
    if prep is not None:
        for col in prep.columns:
            df[t(f"prep.{col}", lang)] = prep[col].values
    return df
